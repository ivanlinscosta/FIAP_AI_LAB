from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
DOCS.mkdir(exist_ok=True)
MAGENTA = 'ED145B'
DARK = '17171C'
GRAY = '66666E'
LIGHT = 'F3F3F5'


def font(size, bold=False, color='000000', name='Arial'):
    return {'name': name, 'size': Pt(size), 'bold': bold, 'color': color}


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def apply_run(run, size=10.5, bold=False, color='000000', font_name='Arial', italic=False):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def setup_doc(title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
    styles = doc.styles
    styles['Normal'].font.name = 'Arial'
    styles['Normal'].font.size = Pt(10.5)
    for sname, size, color in [('Title', 30, 'FFFFFF'), ('Heading 1', 20, DARK), ('Heading 2', 14, MAGENTA), ('Heading 3', 11.5, DARK)]:
        st = styles[sname]
        st.font.name = 'Arial'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    # Custom code style
    if 'Code' not in [s.name for s in styles]:
        st = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = 'Liberation Mono'; st.font.size = Pt(9); st.font.color.rgb = RGBColor.from_string(DARK)
    return doc


def add_cover(doc, title, subtitle):
    sec = doc.sections[0]
    sec.top_margin = Cm(0); sec.bottom_margin = Cm(0); sec.left_margin = Cm(0); sec.right_margin = Cm(0)
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False; t.columns[0].width = Cm(21)
    c = t.cell(0,0); set_cell_shading(c, '09090C'); set_cell_margins(c, 900, 1100, 900, 1100)
    p = c.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run('FIAP  |  AI AGENTS'); apply_run(r, 14, True, MAGENTA)
    p = c.add_paragraph(); p.space_after = Pt(14)
    r = p.add_run(title); apply_run(r, 30, True, 'FFFFFF')
    p = c.add_paragraph(); p.space_after = Pt(20)
    r = p.add_run(subtitle); apply_run(r, 15, False, 'D0D0D5')
    p = c.add_paragraph(); p.space_before = Pt(220)
    r = p.add_run('TOOLS, AUTOMATIONS AND WORKFLOWS'); apply_run(r, 12, True, 'FFFFFF')
    p = c.add_paragraph(); r = p.add_run('Ambiente compartilhado de IA para n8n, Dify e Power Automate'); apply_run(r, 11, False, 'B8B8C0')
    p = c.add_paragraph(); p.space_before = Pt(18)
    r = p.add_run('Versão de laboratório - 2026'); apply_run(r, 10, False, MAGENTA)
    doc.add_page_break()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)


def add_h1(doc, text):
    p = doc.add_paragraph(style='Heading 1'); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style='Heading 2'); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5)
    p.add_run(text); return p


def add_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); apply_run(r, 10.5, True, DARK)
        r = p.add_run(text[len(bold_prefix):]); apply_run(r, 10.5, False, DARK)
    else:
        r = p.add_run(text); apply_run(r, 10.5, False, DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(item); apply_run(r, 10.3, False, DARK)


def add_callout(doc, title, text, fill='F5F5F7', border=MAGENTA):
    table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = table.cell(0,0); set_cell_shading(c, fill); set_cell_margins(c, 160, 180, 160, 180)
    tcPr = c._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left'); left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:color'), border); borders.append(left); tcPr.append(borders)
    p = c.paragraphs[0]; r = p.add_run(title); apply_run(r, 10.5, True, border)
    p = c.add_paragraph(); r = p.add_run(text); apply_run(r, 10.2, False, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = table.cell(0,0); set_cell_shading(c, 'F2F2F4'); set_cell_margins(c, 120, 150, 120, 150)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run('\n')
        r = p.add_run(line); apply_run(r, 9, False, DARK, 'Liberation Mono')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_cell_shading(c, DARK); set_cell_margins(c,120,120,120,120); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(h); apply_run(r,9.5,True,'FFFFFF')
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell_margins(cells[i],110,120,110,120); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0];
            r=p.add_run(str(val)); apply_run(r,9.3,False,DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def architecture_png(path):
    W,H=1600,720
    img=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(img)
    try:
        fb='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'; fr='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
        ftitle=ImageFont.truetype(fb,42); fhead=ImageFont.truetype(fb,28); fbody=ImageFont.truetype(fr,22); fsmall=ImageFont.truetype(fr,19)
    except:
        ftitle=fhead=fbody=fsmall=None
    d.text((60,40),'Arquitetura do FIAP AI Lab',fill='#17171C',font=ftitle)
    def box(x,y,w,h,title,body,fill='#F6F6F8',outline='#C7C7CD'):
        d.rounded_rectangle((x,y,x+w,y+h),24,fill=fill,outline=outline,width=3)
        d.text((x+24,y+24),title,fill='#ED145B',font=fhead)
        yy=y+76
        for line in body:
            d.text((x+24,yy),line,fill='#17171C',font=fbody); yy+=34
    box(90,180,350,230,'OpenAI Project',['Chave real','Somente professor'],fill='#FFF4F8',outline='#ED145B')
    box(625,145,350,300,'LiteLLM Gateway',['Chaves virtuais','Budget por grupo','RPM / TPM','Logs e custos'],fill='#F6F6F8',outline='#17171C')
    box(1160,90,330,150,'n8n',['Workflow / AI nodes'])
    box(1160,285,330,150,'Dify',['LLM / Workflow / RAG'])
    box(1160,480,330,150,'Power Automate',['Custom Connector'])
    # arrows
    for a,b in [((440,295),(625,295)),((975,240),(1160,165)),((975,295),(1160,360)),((975,350),(1160,555))]:
        d.line((a[0],a[1],b[0],b[1]),fill='#ED145B',width=8)
        x,y=b; d.polygon([(x,y),(x-22,y-12),(x-22,y+12)],fill='#ED145B')
    d.text((600,540),'Cada grupo recebe uma chave virtual;\na chave OpenAI nunca sai do servidor.',fill='#55555E',font=fsmall)
    img.save(path)

arch = DOCS / 'arquitetura_fiap_ai_lab.png'
architecture_png(arch)

# -------- Teacher manual --------
doc = setup_doc('Manual do Professor')
add_cover(doc, 'FIAP AI Lab - Manual do Professor', 'Instalação, governança e operação do ambiente compartilhado de IA')
add_h1(doc,'1. Visão geral')
add_p(doc,'O FIAP AI Lab foi desenhado para permitir que grupos de alunos usem n8n, Dify e Power Automate consumindo uma única conta OpenAI do professor, sem revelar a chave real da OpenAI. O componente central é um gateway compatível com a API da OpenAI, responsável por autenticação, chaves virtuais, limites e rastreamento de uso.')
add_callout(doc,'Princípio de segurança','A chave real da OpenAI fica apenas no servidor. Os alunos recebem chaves virtuais revogáveis e com orçamento próprio.')
doc.add_picture(str(arch), width=Cm(16.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_h2(doc,'O que o pacote entrega')
add_bullets(doc,[
    'LiteLLM Proxy como AI Gateway central e compatível com o formato OpenAI.',
    'PostgreSQL para persistência de chaves virtuais e métricas do gateway.',
    'Instância n8n para demonstrações e laboratórios centralizados.',
    'Aliases de modelos: fiap-fast, fiap-standard e fiap-embedding.',
    'Script para criar chaves virtuais por grupo com budget, RPM, TPM e validade.',
    'Workflow n8n de teste, configuração do Dify e OpenAPI 2.0 para Power Automate.',
])

add_h1(doc,'2. Decisões de arquitetura')
add_table(doc,['Camada','Escolha','Motivo'],[
    ['Provedor de IA','OpenAI Project dedicado','Separa consumo da disciplina e facilita governança.'],
    ['Gateway','LiteLLM Proxy','API OpenAI-compatible, chaves virtuais, budgets, rate limiting e tracking.'],
    ['Automação','n8n','Workflows visuais, HTTP, integrações e nodes de IA.'],
    ['AI workflow','Dify','LLM, workflow, RAG e tools em uma interface visual.'],
    ['Microsoft','Power Automate','Integração via Custom Connector quando o tenant/licença permitir.'],
])
add_p(doc,'A arquitetura é intencionalmente híbrida: o gateway é central e controlado pelo professor; as ferramentas dos alunos podem ser hospedadas em uma instância comum ou usadas a partir das contas individuais dos grupos. O objetivo é desacoplar credenciais de provedor das ferramentas de automação.')

add_h1(doc,'3. Preparação do projeto OpenAI')
add_bullets(doc,[
    'Crie um projeto OpenAI exclusivo para a disciplina, separado de usos pessoais ou de produção.',
    'Crie uma chave API dedicada para o laboratório e trate-a como segredo de servidor.',
    'Comece com limites financeiros conservadores e acompanhe o consumo nas primeiras aulas.',
    'Use modelos econômicos como padrão e reserve modelos mais caros para exercícios específicos.',
])
add_callout(doc,'Importante','Não cole a chave OpenAI em workflows exportáveis, screenshots, documentos de aula ou credenciais que os estudantes possam visualizar.')

add_h1(doc,'4. Instalação local ou em VM')
add_p(doc,'Pré-requisitos: Docker Engine e Docker Compose instalados em um computador ou servidor com acesso à internet. Para Dify Cloud e Power Automate, o gateway deve estar exposto por HTTPS público.')
add_h2(doc,'Passo 1 - Gerar o arquivo de ambiente')
add_code(doc,"cp .env.example .env\npython scripts/generate_secrets.py")
add_p(doc,'O script cria automaticamente o master key do LiteLLM, a senha do PostgreSQL e a chave de criptografia do n8n. Depois, abra o arquivo .env e substitua apenas o valor de OPENAI_API_KEY pela chave dedicada da disciplina.')
add_h2(doc,'Passo 2 - Subir os serviços')
add_code(doc,'docker compose up -d')
add_p(doc,'O gateway ficará em http://localhost:4000 e o n8n em http://localhost:5678 no modo local. Em produção, publique o gateway com HTTPS e ajuste GATEWAY_PUBLIC_URL.')
add_h2(doc,'Passo 3 - Criar as chaves dos grupos')
add_code(doc,'python scripts/provision_groups.py')
add_p(doc,'O script cria 10 grupos por padrão. A quantidade, o budget e os limites podem ser alterados no arquivo .env antes da execução.')
add_table(doc,['Parâmetro','Padrão','Uso'],[
    ['GROUP_COUNT','10','Quantidade de chaves/grupos.'],
    ['GROUP_BUDGET_USD','5','Budget máximo por grupo.'],
    ['GROUP_BUDGET_DURATION','30d','Janela de renovação do budget.'],
    ['GROUP_RPM_LIMIT','30','Requisições por minuto.'],
    ['GROUP_TPM_LIMIT','100000','Tokens por minuto.'],
    ['GROUP_KEY_DURATION','45d','Validade da chave virtual.'],
])
add_p(doc,'O arquivo completo secrets/group_keys.csv é exclusivo do professor. Para os alunos, distribua somente o arquivo secrets/grupo-XX.txt correspondente ao grupo.')

add_h1(doc,'5. Como o aluno usa no n8n')
add_h2(doc,'Opção A - Workflow de teste por HTTP')
add_p(doc,'Importe o arquivo n8n/workflows/00_gateway_smoke_test.json. No node FIAP AI Gateway, substitua SEU-GATEWAY pela URL pública do gateway e COLE-A-CHAVE-VIRTUAL-DO-GRUPO pela chave daquele grupo. Execute o workflow.')
add_code(doc,'POST https://SEU-GATEWAY/v1/chat/completions\nAuthorization: Bearer CHAVE_VIRTUAL_DO_GRUPO\nmodel: fiap-fast')
add_h2(doc,'Opção B - OpenAI Chat Model no n8n')
add_p(doc,'No node OpenAI Chat Model, use a chave virtual como API key e informe, nas opções do node, o Base URL do gateway terminado em /v1. Selecione o alias fiap-fast. Essa configuração mantém a experiência de node de IA sem expor a chave real do provedor.')
add_callout(doc,'Didática recomendada','Na Aula 02, comece com HTTP Request para tornar a integração visível. Depois migre para OpenAI Chat Model e mostre que o node apenas abstrai a mesma chamada de API.')

add_h1(doc,'6. Como o aluno usa no Dify')
add_p(doc,'No Dify, instale ou selecione o provider OpenAI-API-compatible. Para cada grupo, configure o endpoint raiz do gateway, a chave virtual e os aliases de modelo liberados.')
add_table(doc,['Campo','Valor'],[
    ['Endpoint URL','https://SEU-GATEWAY (sem /v1)'],
    ['API Key','chave virtual do grupo'],
    ['Model name','fiap-fast'],
    ['Model type','LLM / Chat'],
])
add_p(doc,'Para exercícios de RAG/Knowledge Base, adicione fiap-embedding como modelo de embeddings usando o mesmo endpoint e a mesma chave virtual.')
add_h2(doc,'Teste mínimo')
add_code(doc,'Start -> LLM -> End\nPrompt: Explique em uma frase o que é um workflow automatizado.')

add_h1(doc,'7. Como o aluno usa no Power Automate')
add_p(doc,'A integração está preparada como Custom Connector. Ela depende de o tenant Microsoft permitir conectores personalizados e do licenciamento disponível para os alunos.')
add_h2(doc,'Gerar o Swagger final')
add_code(doc,'python scripts/render_power_automate_swagger.py https://seu-gateway.exemplo.com')
add_p(doc,'Importe power-automate/fiap-ai-gateway.swagger.json em Custom connectors > New custom connector > Import an OpenAPI file. Ao criar a conexão, informe a API key exatamente no formato “Bearer CHAVE_VIRTUAL_DO_GRUPO”.')
add_callout(doc,'Limitação externa','A criação e o compartilhamento de Custom Connectors são controlados pelo ambiente e pela licença Microsoft da FIAP. O pacote deixa a API pronta, mas a instituição precisa permitir o uso no tenant.')

add_h1(doc,'8. Operação em dia de aula')
add_table(doc,['Momento','Ação do professor'],[
    ['Antes da aula','Subir serviços, validar gateway e executar um smoke test.'],
    ['Início','Entregar apenas o cartão de credenciais de cada grupo.'],
    ['Durante','Acompanhar logs, consumo e erros 429/5xx.'],
    ['Se houver abuso','Reduzir budget/rate limit ou revogar a chave do grupo.'],
    ['Final','Exportar evidências, manter ou revogar chaves conforme o plano da turma.'],
])
add_h2(doc,'Smoke test')
add_code(doc,"export FIAP_GROUP_KEY='CHAVE_DO_GRUPO'\npython scripts/smoke_test.py")

add_h1(doc,'9. Governança e controle de custos')
add_bullets(doc,[
    'Uma chave virtual por grupo, nunca uma chave compartilhada para toda a sala.',
    'Budget pequeno por padrão, com ampliação consciente conforme o laboratório exigir.',
    'RPM e TPM limitados para evitar loops acidentais e explosões de custo.',
    'Validade curta das chaves; recrie credenciais a cada turma ou semestre.',
    'Não armazene dados sensíveis dos alunos em prompts de laboratório.',
    'Desative logs de conteúdo quando o exercício usar dados que não devem ser persistidos.',
])
add_h2(doc,'Modelos sugeridos')
add_table(doc,['Alias do laboratório','Modelo atual','Uso didático'],[
    ['fiap-fast','gpt-5-mini','Classificação, extração, resumo, roteamento e testes.'],
    ['fiap-standard','gpt-5.1','Exercícios que exigem raciocínio mais forte.'],
    ['fiap-embedding','text-embedding-3-small','RAG e busca semântica.'],
])
add_p(doc,'Os aliases são intencionais: se você trocar o modelo real no gateway, os workflows dos alunos continuam usando o mesmo nome lógico.')

add_h1(doc,'10. Troubleshooting')
add_table(doc,['Sintoma','Causa provável','Ação'],[
    ['401 Unauthorized','Chave virtual incorreta ou revogada','Confirme o cartão do grupo e o header Bearer.'],
    ['429 Too Many Requests','RPM/TPM ou budget atingido','Aguarde, ajuste o limite ou gere nova chave.'],
    ['Model not found','Alias não liberado para a chave','Use fiap-fast ou revise models na chave virtual.'],
    ['Dify falha no endpoint','URL contém /v1 duplicado','No Dify use a raiz do gateway, sem /v1.'],
    ['n8n não alcança gateway local','Container/host diferente','Use hostname público ou host.docker.internal quando apropriado.'],
    ['Power Automate não importa Swagger','Host inválido ou política do tenant','Gere novamente com URL HTTPS e valide licenciamento/políticas.'],
])

add_h1(doc,'11. O que ainda depende das suas contas')
add_p(doc,'O pacote está configurado e pronto para deploy, mas três ações não podem ser concluídas automaticamente sem acesso às suas contas externas: inserir a chave real do projeto OpenAI, publicar o gateway no provedor de hospedagem escolhido e criar/compartilhar o Custom Connector dentro do tenant Microsoft. Essas ações envolvem credenciais privadas e permissões do proprietário da conta.')
add_callout(doc,'Estado do pacote','Você já possui todos os arquivos, scripts, templates e configurações necessárias. Depois de inserir a chave e publicar o gateway, o provisionamento dos grupos é automático.')

add_h1(doc,'12. Referências técnicas')
refs=[
    'OpenAI API quickstart e boas práticas de chave: https://platform.openai.com/docs/quickstart/make-your-first-api-request',
    'OpenAI models: https://platform.openai.com/docs/models',
    'LiteLLM Proxy / Gateway: https://docs.litellm.ai/',
    'n8n documentation: https://docs.n8n.io/',
    'Dify documentation: https://docs.dify.ai/',
    'Microsoft Custom Connectors: https://learn.microsoft.com/en-us/connectors/custom-connectors/',
]
add_bullets(doc,refs)

teacher_docx = DOCS / 'Manual_do_Professor_FIAP_AI_Lab.docx'
doc.save(teacher_docx)

# -------- Student guide --------
doc = setup_doc('Guia do Aluno')
add_cover(doc, 'FIAP AI Lab - Guia do Aluno', 'Como usar sua credencial de grupo no n8n, Dify e Power Automate')
add_h1(doc,'1. Seu cartão de acesso')
add_p(doc,'Cada grupo receberá um arquivo com Gateway, Base URL, API Key e nomes dos modelos. Essa credencial é exclusiva do grupo e tem limite de uso.')
add_callout(doc,'Não compartilhe sua chave','A chave do grupo pode ser revogada pelo professor. Não publique em GitHub, prints, fóruns ou arquivos compartilhados fora da turma.')
add_h1(doc,'2. n8n')
add_bullets(doc,[
    'Importe o workflow “FIAP - 00 - Teste do AI Gateway”.',
    'Troque SEU-GATEWAY pela URL do cartão do grupo.',
    'Troque COLE-A-CHAVE-VIRTUAL-DO-GRUPO pela sua API Key.',
    'Execute o workflow e confirme a resposta.',
])
add_p(doc,'Para nodes OpenAI Chat Model, use sua chave virtual e configure o Base URL como a “OpenAI-compatible Base URL” do cartão, normalmente terminada em /v1.')
add_h1(doc,'3. Dify')
add_table(doc,['Campo','Preencha com'],[
    ['Provider','OpenAI-API-compatible'],
    ['Endpoint URL','Gateway do cartão, sem /v1'],
    ['API Key','Chave virtual do grupo'],
    ['Model','fiap-fast'],
])
add_p(doc,'Para Knowledge Base/RAG, use fiap-embedding como modelo de embedding.')
add_h1(doc,'4. Power Automate')
add_p(doc,'Se o professor disponibilizar o Custom Connector FIAP AI Gateway no tenant Microsoft, crie uma conexão usando “Bearer SUA-CHAVE” no campo de API key e depois use a ação “Gerar resposta com IA”.')
add_h1(doc,'5. Regras do laboratório')
add_bullets(doc,[
    'Evite loops de execução sem condição de parada.',
    'Não envie dados pessoais ou sensíveis para o modelo.',
    'Use fiap-fast como padrão; utilize modelos avançados somente quando o exercício solicitar.',
    'Se receber erro 429, aguarde antes de tentar novamente e revise o fluxo.',
    'Se uma automação estiver repetindo chamadas, desative o workflow antes de depurar.',
])
add_h1(doc,'6. Diagnóstico rápido')
add_table(doc,['Erro','O que fazer'],[
    ['401','Confira a chave e o prefixo Bearer.'],
    ['429','Você atingiu limite de chamadas ou budget; avise o professor.'],
    ['404/model not found','Use fiap-fast, fiap-standard ou fiap-embedding.'],
    ['Dify sem resposta','Revise endpoint sem /v1 duplicado.'],
    ['n8n repetindo execução','Desative o workflow e procure o trigger/loop.'],
])
student_docx = DOCS / 'Guia_do_Aluno_FIAP_AI_Lab.docx'
doc.save(student_docx)
print(teacher_docx)
print(student_docx)
